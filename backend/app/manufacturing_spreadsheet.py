"""A replaceable, linked directory report, independent of recovery archives."""

import hashlib
import io
import os
import time
from collections import defaultdict
from uuid import uuid4

from fastapi import APIRouter, Request, Response
from google.cloud import firestore
from googleapiclient.http import MediaIoBaseUpload
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

from .domain import Identifier, Model, Problem, now
from .manufacturing_backups import FILES, ManufacturingDriveBackupStore, checked_data
from .storage import digest, encode

KINDS = ("products", "operations", "suppliers", "contacts", "steps", "assignments")
NAME = "ProductTrackr Directory.xlsx"
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PATH = "/api/manufacturing/settings/spreadsheet"
MAX_BYTES = 32 * 1024 * 1024
MAX_CELLS = 200_000


def workbook(data, identity, model_names=None):
    model_names = model_names or {}
    if any(model_id not in model_names for product in data["products"].values()
           for model_id in product.get("modelIds", [])):
        raise Problem(
            "A linked machine model is unavailable. Refresh before exporting.", 409, "model_missing",
        )
    # Reuse the graph validator without reading or exporting execution collections.
    try:
        checked_data(
            {
                **{kind: data[kind] if kind in KINDS else {} for kind in FILES},
                "settings": {"weeklyEnabled": True},
                **({"productFolders": data["productFolders"]} if "productFolders" in data else {}),
            },
            directories_only=True,
        )
    except Problem:
        raise Problem(
            "The directory relationships are inconsistent. No spreadsheet was replaced.",
            422,
            "spreadsheet_data",
        ) from None
    ordered = {
        kind: sorted(data[kind].values(), key=lambda row: (row["name"].casefold(), row["id"]))
        for kind in ("products", "operations", "suppliers", "contacts")
    }
    ordered["contacts"].sort(
        key=lambda row: (
            data["suppliers"][row["supplierId"]]["name"].casefold(),
            row["supplierId"],
            row["name"].casefold(),
            row["id"],
        )
    )
    steps = [
        (product, position, data["steps"][key])
        for product in ordered["products"]
        for position, key in enumerate(product["stepIds"], 1)
    ]
    positions = {step["id"]: position for _, position, step in steps}
    step_rank = {step["id"]: index for index, (_, _, step) in enumerate(steps)}
    assignments = sorted(
        data["assignments"].values(),
        key=lambda row: (
            step_rank[row["stepId"]],
            data["suppliers"][row["supplierId"]]["name"].casefold(),
            row["id"],
        ),
    )
    row_numbers = {
        title: {row["id"]: index for index, row in enumerate(rows, 2)}
        for title, rows in (
            ("Products", ordered["products"]),
            ("Operations", ordered["operations"]),
            ("Suppliers", ordered["suppliers"]),
            ("Contacts", ordered["contacts"]),
            ("Product operations", [step for _, _, step in steps]),
            ("Supplier links", assignments),
        )
    }
    links, contacts = defaultdict(list), defaultdict(list)
    for row in assignments:
        links[row["stepId"]].append(row)
    for row in ordered["contacts"]:
        contacts[row["supplierId"]].append(row)

    def ref(label, sheet, key):
        return (label, f"'{sheet}'!A{row_numbers[sheet][key]}") if key else label

    book = Workbook()
    book.remove(book.active)
    cell_count = 0

    def sheet(title, headers, rows):
        nonlocal cell_count
        page = book.create_sheet(title)
        page.freeze_panes = "A2"
        page.sheet_properties.pageSetUpPr.fitToPage = True
        for index, values in enumerate([headers, *rows], 1):
            cell_count += len(values)
            if cell_count > MAX_CELLS:
                raise Problem("The directory report exceeds 200,000 cells.", 413, "spreadsheet_size")
            for column, value in enumerate(values, 1):
                label, target = value if isinstance(value, tuple) else (value, None)
                if isinstance(label, str):
                    if (
                        len(label) > 32767
                        or ILLEGAL_CHARACTERS_RE.search(label)
                        or any(
                            0xD800 <= ord(char) <= 0xDFFF or ord(char) in (0xFFFE, 0xFFFF) for char in label
                        )
                    ):
                        raise Problem(
                            "A directory value cannot be represented safely in a spreadsheet. "
                            "No existing spreadsheet was replaced.",
                            422,
                            "spreadsheet_value",
                        )
                cell = page.cell(index, column, label)
                if isinstance(label, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if target:
                    cell.hyperlink = (
                        Hyperlink(ref=cell.coordinate, target=target)
                        if target.startswith("https://")
                        else Hyperlink(ref=cell.coordinate, location=target)
                    )
                    cell.font = Font(color="245B85", underline="single")
                if index == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1C354C")
        page.auto_filter.ref = page.dimensions
        for column in page.columns:
            width = max(len(str(cell.value or "").split("\n")[0]) for cell in column)
            dimension = page.column_dimensions[column[0].column_letter]
            dimension.width = min(48, max(16, width + 2))
            dimension.hidden = str(column[0].value).endswith(" key")
        page.row_dimensions[1].height = 30

    sheet(
        "Products",
        [
            "Part number", "Product", "Operations (in order)", "Process",
            "Models", "Stock quantity", "Minimum stock quantity", "Notes", "Drive folder", "Product key",
        ],
        [
            [
                product["partNumber"],
                product["name"],
                "\n".join(
                    f"{index}. {data['operations'][data['steps'][key]['operationId']]['name']}"
                    for index, key in enumerate(product["stepIds"], 1)
                ),
                ref("View process", "Product operations", product["stepIds"][0])
                if product["stepIds"]
                else "No operations",
                ", ".join(model_names[key] for key in product["modelIds"])
                if product.get("modelIds") else "Common",
                product.get("stockQuantity"),
                product.get("minimumStockQuantity"),
                product["notes"],
                (folder["url"], folder["url"]) if folder else "",
                product["id"],
            ]
            for product in ordered["products"]
            for folder in [data.get("productFolders", {}).get("folders", {}).get(product["id"])]
        ],
    )
    sheet(
        "Product operations",
        [
            "Part number",
            "Product",
            "Step",
            "Operation",
            "Available suppliers",
            "Instructions",
            "Primary supplier",
            "Primary contact",
            "Supplier options",
            "Supplier links",
            "Product key",
            "Step key",
            "Operation key",
            "Primary link key",
        ],
        [
            [
                product["partNumber"],
                ref(product["name"], "Products", product["id"]),
                position,
                ref(data["operations"][step["operationId"]]["name"], "Operations", step["operationId"]),
                "\n".join(
                    data["suppliers"][link["supplierId"]]["name"]
                    + (" (primary)" if link["id"] == step["primaryAssignmentId"] else "")
                    for link in links[step["id"]]
                ),
                step["instructions"],
                ref(data["suppliers"][primary["supplierId"]]["name"], "Suppliers", primary["supplierId"])
                if primary
                else "",
                ref(data["contacts"][primary["contactId"]]["name"], "Contacts", primary["contactId"])
                if primary and primary["contactId"]
                else "",
                len(links[step["id"]]),
                ref("View supplier options", "Supplier links", links[step["id"]][0]["id"])
                if links[step["id"]]
                else "No linked suppliers",
                product["id"],
                step["id"],
                step["operationId"],
                step["primaryAssignmentId"] or "",
            ]
            for product, position, step in steps
            for primary in [data["assignments"].get(step["primaryAssignmentId"])]
        ],
    )
    sheet(
        "Supplier links",
        [
            "Part number",
            "Product",
            "Step",
            "Operation",
            "Supplier",
            "Supplier type",
            "Relationship",
            "Contact",
            "Primary",
            "Link notes",
            "Product key",
            "Step key",
            "Operation key",
            "Supplier key",
            "Contact key",
            "Link key",
        ],
        [
            [
                product["partNumber"],
                ref(product["name"], "Products", product["id"]),
                ref(positions[step["id"]], "Product operations", step["id"]),
                ref(data["operations"][step["operationId"]]["name"], "Operations", step["operationId"]),
                ref(supplier["name"], "Suppliers", supplier["id"]),
                supplier["type"],
                supplier["relationship"],
                ref(data["contacts"][link["contactId"]]["name"], "Contacts", link["contactId"])
                if link["contactId"]
                else "",
                "Yes" if step["primaryAssignmentId"] == link["id"] else "No",
                link["notes"],
                product["id"],
                step["id"],
                step["operationId"],
                supplier["id"],
                link["contactId"] or "",
                link["id"],
            ]
            for link in assignments
            for product, step, supplier in [
                (
                    data["products"][link["productId"]],
                    data["steps"][link["stepId"]],
                    data["suppliers"][link["supplierId"]],
                )
            ]
        ],
    )
    sheet(
        "Operations",
        ["Operation", "Description", "Operation key"],
        [[row["name"], row["description"], row["id"]] for row in ordered["operations"]],
    )
    sheet(
        "Suppliers",
        [
            "Supplier",
            "Type",
            "Relationship",
            "Address",
            "Email",
            "Phone",
            "Notes",
            "Contacts",
            "Supplier key",
        ],
        [
            [
                row["name"],
                row["type"],
                row["relationship"],
                row["address"],
                row["email"],
                row["phone"],
                row["notes"],
                ref(f"View contacts ({len(contacts[row['id']])})", "Contacts", contacts[row["id"]][0]["id"])
                if contacts[row["id"]]
                else "No company contacts",
                row["id"],
            ]
            for row in ordered["suppliers"]
        ],
    )
    sheet(
        "Contacts",
        ["Contact", "Company", "Designation", "Email", "Phone", "Contact key", "Supplier key"],
        [
            [
                row["name"],
                ref(data["suppliers"][row["supplierId"]]["name"], "Suppliers", row["supplierId"]),
                row["designation"],
                row["email"],
                row["phone"],
                row["id"],
                row["supplierId"],
            ]
            for row in ordered["contacts"]
        ],
    )
    sheet(
        "Export info",
        ["Field", "Value"],
        [
            ["Exported at (UTC)", identity["exportedAt"]],
            ["Workspace", identity["workspaceId"]],
            ["Source generation", identity["generation"]],
            ["Source sequence", identity["sourceSequence"]],
            ["Contents", "All directory entities, ordered Product operations and supplier/contact links."],
            [
                "Relationships",
                "Keys are stable internal IDs; the user-entered Part number is not a unique key.",
            ],
            [
                "Usage",
                "Read-only export. Changes here do not sync back. The next export replaces the workbook.",
            ],
            [
                "Technical columns",
                "Relationship-key columns are retained but hidden. Unhide them when needed.",
            ],
            [
                "Not included",
                "Manufacturing jobs, history, counters, receipts and backups. Use ZIP backups for recovery.",
            ],
            *[[kind, len(data[kind])] for kind in KINDS],
        ],
    )
    output = io.BytesIO()
    book.save(output)
    if output.tell() > MAX_BYTES:
        raise Problem("The directory report exceeds 32 MiB.", 413, "spreadsheet_size")
    return output.getvalue()


class LocalSpreadsheetStore:
    provider = "local"

    def __init__(self, settings, scope):
        self.path = settings.data_dir / "producttrackr-spreadsheets" / scope / "ProductTrackr Directory.xlsx"

    def allocate(self):
        return "directory-workbook"

    def read(self, file_id):
        if self.path.exists() and self.path.stat().st_size > MAX_BYTES:
            raise Problem("The directory workbook exceeds 32 MiB.", 413, "spreadsheet_size")
        return self.path.read_bytes() if self.path.exists() else None

    def put(self, file_id, data, *, existing=False, before_write=None):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.path.with_suffix(f".{uuid4().hex}.tmp")
        try:
            temporary.write_bytes(data)
            if before_write:
                before_write()
            os.replace(temporary, self.path)
        finally:
            temporary.unlink(missing_ok=True)


class DriveSpreadsheetStore(ManufacturingDriveBackupStore):
    provider = "drive"
    app = "producttrackr-directory-spreadsheet"

    def __init__(self, settings, scope):
        super().__init__(settings)
        self.scope = scope

    def _folder(self):
        return self._root_folder()

    def allocate(self):
        folder = self._folder()
        matches = list(
            self.backend._list(
                f"'{folder}' in parents and trashed = false "
                f"and appProperties has {{ key='app' and value='{self.app}' }} "
                f"and appProperties has {{ key='workspace' and value='{self.scope}' }}",
                fields="id",
            )
        )
        if len(matches) > 1:
            raise Problem("More than one managed directory workbook exists.", 409, "spreadsheet_ambiguous")
        return matches[0]["id"] if matches else self.backend.allocate()

    def _metadata(self, file_id):
        folder = self._folder()
        try:
            row = self.backend._execute(
                self.backend._api()
                .files()
                .get(
                    fileId=file_id,
                    fields="id,parents,mimeType,trashed,appProperties,size,md5Checksum",
                )
            )
        except Problem as error:
            if error.code == "drive_missing":
                return None
            raise
        if (
            row.get("trashed")
            or row.get("parents") != [folder]
            or row.get("mimeType") != MIME
            or row.get("appProperties", {}).get("app") != self.app
            or row.get("appProperties", {}).get("workspace") != self.scope
        ):
            raise Problem(
                "The managed workbook was moved or replaced. Restore it to the configured folder.",
                409,
                "spreadsheet_changed",
            )
        if int(row.get("size", MAX_BYTES + 1)) > MAX_BYTES:
            raise Problem("The directory workbook exceeds 32 MiB.", 413, "spreadsheet_size")
        token = None
        while True:
            permissions = self.backend._execute(
                self.backend._api()
                .permissions()
                .list(
                    fileId=file_id,
                    fields="nextPageToken,permissions(type)",
                    pageSize=100,
                    pageToken=token,
                )
            )
            if any(item["type"] in ("anyone", "domain") for item in permissions["permissions"]):
                raise Problem("The managed workbook must remain private.", 503, "spreadsheet_permissions")
            token = permissions.get("nextPageToken")
            if not token:
                break
        return row

    def put(self, file_id, data, *, existing=False, before_write=None):
        # Capture the ETag before ownership checks so concurrent moves or edits invalidate the write.
        try:
            etag = self.backend.etag(file_id)
        except Problem as error:
            if error.code != "drive_missing":
                raise
            etag = None
        metadata = self._metadata(file_id)
        media = MediaIoBaseUpload(io.BytesIO(data), mimetype=MIME, resumable=False)
        files = self.backend._api().files()
        if metadata:
            if not etag:
                raise Problem("The workbook changed. Retry export.", 409, "spreadsheet_changed")
            request = (
                self.backend._conditional_api()
                .files()
                .update(
                    fileId=file_id,
                    body={"title": NAME},
                    media_body=media,
                    fields="id",
                )
            )
            request.headers["If-Match"] = etag
        else:
            if existing:
                raise Problem(
                    "The managed workbook is missing. Restore its original Drive file before exporting.",
                    409,
                    "spreadsheet_changed",
                )
            request = files.create(
                body={
                    "id": file_id,
                    "name": NAME,
                    "mimeType": MIME,
                    "parents": [self._folder()],
                    "appProperties": {"app": self.app, "workspace": self.scope},
                },
                media_body=media,
                fields="id",
            )
        if before_write:
            before_write()
        self.backend._execute(request)


class SpreadsheetExport:
    def __init__(self, repository, settings, store=None):
        self.repo, self.settings = repository, settings
        scope = digest(settings.manufacturing_identity(self.repo.workspace.id))[:24]
        self.store = store or (
            DriveSpreadsheetStore(settings, scope)
            if settings.manufacturing_backup_storage == "drive"
            else LocalSpreadsheetStore(settings, scope)
        )

    def admin(self, actor):
        return (
            self.settings.mode == "demo"
            or actor.casefold() in (self.settings.manufacturing_admins or self.settings.recovery_admins)
        )

    def require_admin(self, actor):
        if not self.admin(actor):
            raise Problem("Spreadsheet export requires a workspace administrator.", 403, "export_admin")

    def status(self, actor):
        control = self.repo.workspace.get().to_dict()
        state = control.get("directorySpreadsheet") or {}
        matching = state.get("provider") == self.store.provider
        return {
            **self.repo.meta(),
            "provider": self.store.provider,
            "canExport": self.admin(actor),
            "name": NAME,
            "folderUrl": (
                f"https://drive.google.com/drive/folders/{self.settings.root_id}"
                if self.settings.root_id
                else None
            ),
            "spreadsheet": state.get("result") if matching and state.get("completed") else None,
            "lastError": state.get("lastError") if matching else None,
            "pendingWriteId": state.get("writeId")
            if (
                matching
                and not state.get("completed")
                and state.get("actor") == actor
                and state.get("generation") == control["generation"]
            )
            else None,
        }

    def _guard(self, control, owner):
        lease = control.get("maintenance") or {}
        if (
            control["generation"] != self.repo.control["generation"]
            or control["activeDataset"] != self.repo.control["activeDataset"]
            or lease.get("owner") != owner
            or lease.get("expires", 0) <= time.time()
        ):
            raise Problem("Export ownership changed. Retry after reloading.", 409, "workspace_changed")

    def _update(self, owner, fields, *, finish=False):
        @firestore.transactional
        def update(transaction):
            control = self.repo.workspace.get(transaction=transaction).to_dict()
            self._guard(control, owner)
            transaction.update(
                self.repo.workspace,
                {
                    "directorySpreadsheet": {**control["directorySpreadsheet"], **fields},
                    "maintenance": None
                    if finish
                    else {**control["maintenance"], "expires": time.time() + 600},
                },
            )

        update(self.repo.client.transaction())

    def _release(self, owner, error):
        @firestore.transactional
        def release(transaction):
            control = self.repo.workspace.get(transaction=transaction).to_dict()
            if (control.get("maintenance") or {}).get("owner") == owner:
                transaction.update(
                    self.repo.workspace,
                    {
                        "maintenance": None,
                        "directorySpreadsheet.lastError": error,
                    },
                )

        release(self.repo.client.transaction())

    def export(self, write_id, actor, workspace, generation):
        self.require_admin(actor)
        owner = uuid4().hex

        @firestore.transactional
        def claim(transaction):
            control = self.repo.workspace.get(transaction=transaction).to_dict()
            if (
                workspace != self.repo.workspace.id
                or generation != control["generation"]
                or generation != self.repo.control["generation"]
            ):
                raise Problem("ProductTrackr changed. Reload before exporting.", 409, "workspace_changed")
            state = control.get("directorySpreadsheet") or {}
            if state.get("provider") != self.store.provider:
                state = {}
            if state.get("writeId") == write_id:
                if state.get("actor") != actor:
                    raise Problem("Export identifier belongs to another request.", 409, "write_id_reused")
                if state.get("completed") and state.get("generation") == generation:
                    return state, True
            if control.get("pendingRecordCleanup"):
                raise Problem("Resume record cleanup before exporting.", 503, "cleanup_pending")
            maintenance = control.get("maintenance")
            if maintenance and maintenance.get("expires", float("inf")) > time.time():
                raise Problem("Storage maintenance is in progress. Retry shortly.", 503, "maintenance")
            if (control.get("backupLease") or {}).get("expires", 0) > time.time():
                raise Problem("A backup is in progress. Retry shortly.", 503, "backup_in_progress")
            if (
                self.store.provider == "drive"
                and state.get("fileId")
                and state.get("folderId") != self.settings.root_id
            ):
                raise Problem("The configured spreadsheet folder changed.", 409, "spreadsheet_changed")
            state = {
                **state,
                "provider": self.store.provider,
                "writeId": write_id,
                "actor": actor,
                "generation": generation,
                "completed": False,
                "lastError": None,
                "folderId": self.settings.root_id,
            }
            transaction.update(
                self.repo.workspace,
                {
                    "directorySpreadsheet": state,
                    "maintenance": {
                        "writeId": f"spreadsheet-{write_id}",
                        "owner": owner,
                        "expires": time.time() + 600,
                    },
                },
            )
            return state, False

        state, completed = claim(self.repo.client.transaction())
        if completed:
            return state["result"]
        try:
            file_id = state.get("fileId") or self.store.allocate()
            self._update(owner, {"fileId": file_id})
            data, size = {}, 0
            for kind in KINDS:
                data[kind] = {}
                for row in self.repo.collection(kind).stream():
                    value = row.to_dict()
                    size += len(encode(value))
                    if size > MAX_BYTES:
                        raise Problem("The directory snapshot exceeds 32 MiB.", 413, "spreadsheet_size")
                    data[kind][row.id] = value
            data["productFolders"] = self.repo.product_folders(data["products"])
            control = self.repo.workspace.get().to_dict()
            self._guard(control, owner)
            identity = {
                "workspaceId": workspace,
                "generation": generation,
                "sourceSequence": control["sequence"],
                "exportedAt": now(),
            }
            from .firestore_storage import FirestoreRepository

            sales = FirestoreRepository(self.settings, client=self.repo.client)
            before = sales.workspace.get().to_dict()
            model_names = {row["entityId"]: row["data"]["name"] for row in sales.list("model")}
            after = sales.workspace.get().to_dict()
            if any(before[key] != after[key] for key in ("generation", "sequence")):
                raise Problem("Models changed during export. Retry the export.", 409, "models_changed")
            identity.update(salesGeneration=before["generation"], salesSequence=before["sequence"])
            content = workbook(data, identity, model_names)
            with self.repo.lock:
                self._update(owner, {})
                self.store.put(
                    file_id,
                    content,
                    existing=bool(state.get("result")),
                    before_write=lambda: self._update(owner, {}),
                )
                if self.store.read(file_id) != content:
                    raise Problem("Spreadsheet readback failed. Retry export.", 503, "spreadsheet_verify")
            result = {
                **identity,
                "id": file_id,
                "name": NAME,
                "provider": self.store.provider,
                "counts": {kind: len(data[kind]) for kind in KINDS},
                "sha256": hashlib.sha256(content).hexdigest(),
                "sizeBytes": len(content),
                "downloadUrl": f"{PATH}/download",
                "url": (
                    f"https://drive.google.com/file/d/{file_id}/view"
                    if self.store.provider == "drive"
                    else None
                ),
            }
            self._update(owner, {"completed": True, "result": result}, finish=True)
            return result
        except Exception as error:
            message = (
                error.message if isinstance(error, Problem) else "Export was not confirmed. Retry export."
            )
            self._release(owner, message)
            if isinstance(error, OSError):
                raise Problem(message, 503, "spreadsheet_storage") from None
            raise

    def download(self, actor):
        self.require_admin(actor)
        state = self.repo.workspace.get().to_dict().get("directorySpreadsheet") or {}
        if not state.get("completed") or state.get("provider") != self.store.provider:
            raise Problem("Create a verified spreadsheet export first.", 404, "spreadsheet_missing")
        content = self.store.read(state["fileId"])
        if content is None or hashlib.sha256(content).hexdigest() != state["result"]["sha256"]:
            raise Problem("The workbook changed or is missing. Export again.", 409, "spreadsheet_changed")
        return content


class ExportWrite(Model):
    writeId: Identifier


def manufacturing_spreadsheet_router(repository, settings):
    router = APIRouter()

    def manager():
        return SpreadsheetExport(repository(), settings)

    @router.get("/settings/spreadsheet")
    def status(request: Request):
        return manager().status(request.state.actor)

    @router.post("/settings/spreadsheet")
    def export(body: ExportWrite, request: Request):
        return manager().export(
            body.writeId,
            request.state.actor,
            request.headers.get("x-producttrackr-workspace"),
            request.headers.get("x-producttrackr-generation"),
        )

    @router.get("/settings/spreadsheet/download")
    def download(request: Request):
        return Response(
            manager().download(request.state.actor),
            media_type=MIME,
            headers={"Content-Disposition": f'attachment; filename="{NAME}"', "Cache-Control": "no-store"},
        )

    return router
