import io
import time
from dataclasses import replace
from unittest.mock import Mock, patch
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.domain import Problem
from backend.app.main import create_app
from backend.app.manufacturing_backups import ManufacturingBackupManager
from backend.app.manufacturing_spreadsheet import (
    MIME,
    NAME,
    PATH,
    DriveSpreadsheetStore,
    LocalSpreadsheetStore,
    SpreadsheetExport,
)
from backend.app.manufacturing_storage import ManufacturingRepository
from backend.tests.test_manufacturing import app as app
from backend.tests.test_manufacturing import command, create, detail
from backend.tests.test_manufacturing_backups import backup, headers, inspect, restore
from backend.tests.test_manufacturing_records import job
from backend.tests.test_manufacturing_records import process as process


def export(client, write_id=None, expected=200, request_headers=None):
    response = client.post(
        PATH,
        json={"writeId": write_id or uuid4().hex},
        headers=request_headers or headers(client),
    )
    assert response.status_code == expected, response.text
    return response.json()


def download(client):
    response = client.get(f"{PATH}/download")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == MIME
    return load_workbook(io.BytesIO(response.content)), response.content


def rows(book, title):
    sheet = list(book[title].rows)
    return [{head.value: cell.value for head, cell in zip(sheet[0], row)} for row in sheet[1:]]


def unchanged(repo, before):
    data, control = repo.snapshot()
    assert data == before[0]
    assert all(
        control[key] == before[1][key] for key in ("generation", "version", "sequence", "activeDataset")
    )
    assert not control.get("maintenance")


def test_workbook_preserves_every_directory_relationship_and_literal_value(process):
    client, settings, product, company, contact, other_company, other_contact, individual = process
    unused = create(client, "products", name="Unused product", partNumber=product["partNumber"])
    operation = create(client, "operations", name="Unused operation", description='=HYPERLINK("bad","bad")')
    lone = create(client, "suppliers", name="Standalone supplier", phone="0012345678")
    unlinked = create(
        client, "contacts", name="Unlinked contact", supplierId=company["id"], phone="+91 01234",
    )
    current = detail(client, product)["entity"]
    command(
        client,
        f"/products/{product['id']}/reorder",
        {"stepIds": current["stepIds"][::-1]},
        current["version"],
    )
    job(client, product)
    repo = ManufacturingRepository(settings)
    before, identity = repo.snapshot()
    sales = client.get("/api/accounts").json()["items"]
    result = export(client)
    book, content = download(client)
    assert result["provider"] == "local" and result["url"] is None
    assert result["counts"] == {kind: len(before[kind]) for kind in result["counts"]}
    assert book.sheetnames == [
        "Products",
        "Product operations",
        "Supplier links",
        "Operations",
        "Suppliers",
        "Contacts",
        "Export info",
    ]
    products = {row["Product key"]: row for row in rows(book, "Products")}
    assert set(products) == {product["id"], unused["id"]}
    assert products[product["id"]]["Part number"] == product["partNumber"]
    assert products[unused["id"]]["Process"] == "No operations"
    assert products[product["id"]]["Operations (in order)"] == "1. Machining\n2. Machining\n3. Machining"
    route = rows(book, "Product operations")
    assert [row["Step key"] for row in route] == before["products"][product["id"]]["stepIds"]
    assert [row["Step"] for row in route] == [1, 2, 3]
    assert all(
        row["Primary supplier"] == company["name"]
        and row["Primary contact"] == contact["name"]
        and row["Supplier options"] == 2
        for row in route
    )
    assert all(
        f"{company['name']} (primary)" in row["Available suppliers"]
        and individual["name"] in row["Available suppliers"]
        for row in route
    )
    supplier_links = rows(book, "Supplier links")
    assert {row["Link key"] for row in supplier_links} == before["assignments"].keys()
    for row in supplier_links:
        original = before["assignments"][row["Link key"]]
        assert row["Product key"] == original["productId"] and row["Step key"] == original["stepId"]
        assert row["Supplier key"] == original["supplierId"]
        assert (row["Contact key"] or None) == original["contactId"]
        assert row["Primary"] == ("Yes" if original["supplierId"] == company["id"] else "No")
    assert {row["Contact key"] for row in rows(book, "Contacts")} == {
        contact["id"],
        other_contact["id"],
        unlinked["id"],
    }
    suppliers = {row["Supplier key"]: row for row in rows(book, "Suppliers")}
    assert suppliers[other_company["id"]]["Relationship"] == "internal"
    assert suppliers[individual["id"]]["Type"] == "individual"
    assert suppliers[lone["id"]]["Phone"] == "0012345678"
    operations = rows(book, "Operations")
    index = next(index for index, row in enumerate(operations, 2) if row["Operation key"] == operation["id"])
    assert book["Operations"].cell(index, 2).value == '=HYPERLINK("bad","bad")'
    assert book["Operations"].cell(index, 2).data_type == "s"
    for sheet in book:
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref
        for column in sheet[1]:
            assert sheet.column_dimensions[column.column_letter].hidden == str(column.value).endswith(" key")
        for row in sheet:
            for cell in row:
                assert cell.data_type != "f"
                if cell.hyperlink:
                    target, coordinate = cell.hyperlink.location.split("!")
                    assert book[target.strip("'")][coordinate].row <= book[target.strip("'")].max_row
    assert b"manufacturing-records.json" not in content
    unchanged(repo, (before, identity))
    assert client.get("/api/accounts").json()["items"] == sales
    assert ManufacturingBackupManager(repo, settings).rows() == []
    assert client.get(PATH).json()["spreadsheet"] == result


def test_updates_same_workbook_without_stale_rows_or_growing_metadata(process):
    client, settings, product, *_ = process
    first = export(client)
    before_files = list((settings.data_dir / "producttrackr-spreadsheets").rglob("*.xlsx"))
    assert len(before_files) == 1
    current = detail(client, product)
    non_primary = next(row for row in current["assignments"] if row["contactId"] is None)
    step = next(row for row in current["steps"] if row["id"] == non_primary["stepId"])
    response, _ = command(
        client,
        f"/steps/{step['id']}/assignments/{non_primary['id']}/remove",
        {},
        step["version"],
        non_primary["id"],
    )
    assert response.status_code == 200, response.text
    product = detail(client, product)["entity"]
    response, _ = command(
        client,
        f"/products/{product['id']}",
        {"name": "Renamed product", "partNumber": product["partNumber"], "notes": "Updated report"},
        product["version"],
        product["id"],
        "PUT",
    )
    assert response.status_code == 200
    write_id = uuid4().hex
    second = export(client, write_id)
    assert second["id"] == first["id"]
    assert second["counts"]["assignments"] == first["counts"]["assignments"] - 1
    assert export(client, write_id) == second
    book, _ = download(client)
    assert rows(book, "Products")[0]["Product"] == "Renamed product"
    assert rows(book, "Products")[0]["Notes"] == "Updated report"
    assert non_primary["id"] not in {row["Link key"] for row in rows(book, "Supplier links")}
    assert list((settings.data_dir / "producttrackr-spreadsheets").rglob("*.xlsx")) == before_files
    assert set(ManufacturingRepository(settings).workspace.get().to_dict()["directorySpreadsheet"]) == {
        "provider",
        "writeId",
        "actor",
        "generation",
        "completed",
        "lastError",
        "folderId",
        "fileId",
        "result",
    }


def test_empty_workspace_exports_headers_without_seeding_entities(app):
    client, settings = app
    result = export(client)
    assert set(result["counts"].values()) == {0}
    book, _ = download(client)
    assert all(book[title].max_row == 1 for title in book.sheetnames[:-1])
    assert not ManufacturingRepository(settings).snapshot()[0]["products"]


def test_products_sheet_contains_clickable_drive_links_before_and_after_restore(app):
    _, settings = app
    configured = replace(settings, root_id="manufacturing-root")
    with TestClient(create_app(configured)) as client:
        product = create(client, "products", name="Product with drawings", partNumber="PART-042")
        unconnected = create(client, "products", name="Product without a folder")
        repo = ManufacturingRepository(configured)
        reference = repo.workspace.collection("productFolders").document(product["id"])
        reference.set({"id": "original-product-folder", "parentId": "manufacturing-products"})
        expected = "https://drive.google.com/drive/folders/original-product-folder"

        def check_links():
            with patch("backend.app.manufacturing_files.ManufacturingDriveBackupStore") as drive:
                export(client)
                drive.assert_not_called()
            book, _ = download(client)
            products = {row["Product key"]: row for row in rows(book, "Products")}
            assert products[product["id"]]["Drive folder"] == expected
            assert products[unconnected["id"]]["Drive folder"] is None
            page = book["Products"]
            column = next(cell.column for cell in page[1] if cell.value == "Drive folder")
            row = next(
                index for index, value in enumerate(rows(book, "Products"), 2)
                if value["Product key"] == product["id"]
            )
            assert page.cell(row, column).hyperlink.target == expected
            assert page.cell(row, column).hyperlink.location is None
            assert page.cell(row, column).data_type == "s"

        check_links()
        saved = backup(client)
        reference.delete()
        restore(client, saved, inspect(client, saved))
        check_links()
        assert not reference.get().exists


def test_authorization_csrf_and_workspace_fencing_precede_export(process):
    client, settings, *_ = process
    assert client.get(f"{PATH}/download").status_code == 404
    export(client, expected=409, request_headers={
        **headers(client), "X-ProductTrackr-Generation": uuid4().hex,
    })
    export(client, expected=409, request_headers={
        **headers(client), "X-ProductTrackr-Workspace": uuid4().hex,
    })
    assert client.post(PATH, json={"writeId": uuid4().hex}).status_code == 403
    with patch.object(SpreadsheetExport, "admin", return_value=False):
        assert client.get(PATH).json()["canExport"] is False
        export(client, expected=403)
        assert client.get(f"{PATH}/download").status_code == 403
    assert not list((settings.data_dir / "producttrackr-spreadsheets").rglob("*.xlsx"))


def test_failed_upload_and_unverified_readback_retry_same_file(process):
    client, settings, *_ = process
    repo = ManufacturingRepository(settings)
    before = repo.snapshot()
    write_id = uuid4().hex
    with patch.object(LocalSpreadsheetStore, "put", side_effect=Problem("Upload failed", 503)):
        export(client, write_id, 503)
    unchanged(repo, before)
    state = client.get(PATH).json()
    assert state["pendingWriteId"] == write_id and state["lastError"] == "Upload failed"
    file_id = repo.workspace.get().to_dict()["directorySpreadsheet"]["fileId"]
    with patch.object(LocalSpreadsheetStore, "read", return_value=b"unverified"):
        export(client, write_id, 503)
    assert client.get(PATH).json()["spreadsheet"] is None
    assert not repo.workspace.get().to_dict().get("maintenance")
    result = export(client, write_id)
    assert result["id"] == file_id and export(client, write_id) == result
    assert len(list((settings.data_dir / "producttrackr-spreadsheets").rglob("*.xlsx"))) == 1
    unchanged(repo, before)
    download(client)


@pytest.mark.parametrize(
    "state",
    [
        {"pendingRecordCleanup": "pending"},
        {"maintenance": {"owner": "other", "expires": time.time() + 600}},
        {"backupLease": {"owner": "other", "expires": time.time() + 600}},
    ],
)
def test_export_cannot_cross_maintenance_or_cleanup(process, state):
    client, settings, *_ = process
    repo = ManufacturingRepository(settings)
    repo.workspace.update(state)
    export(client, expected=503)
    assert not list((settings.data_dir / "producttrackr-spreadsheets").rglob("*.xlsx"))


def test_export_fences_mutations_and_backups_until_verified(process):
    client, _, *_ = process
    original = LocalSpreadsheetStore.put

    def during_export(store, *args, **kwargs):
        assert command(client, "/products", {"name": "Cannot edit during export"})[0].status_code == 503
        export(client, expected=503)
        assert (
            client.post(
                "/api/backups",
                json={"writeId": uuid4().hex},
                headers=headers(client),
            ).status_code
            == 503
        )
        return original(store, *args, **kwargs)

    with patch.object(LocalSpreadsheetStore, "put", during_export):
        export(client)
    assert command(client, "/products", {"name": "Can edit after export"})[0].status_code == 200


def test_generation_takeover_prevents_stale_worker_upload(process):
    client, settings, *_ = process
    repo = ManufacturingRepository(settings)
    original = SpreadsheetExport._update

    def takeover(manager, owner, fields, **kwargs):
        if not fields:
            repo.workspace.update({"generation": uuid4().hex, "maintenance": None})
        return original(manager, owner, fields, **kwargs)

    with patch.object(SpreadsheetExport, "_update", takeover):
        export(client, expected=409)
    assert not list((settings.data_dir / "producttrackr-spreadsheets").rglob("*.xlsx"))
    assert not repo.workspace.get().to_dict().get("maintenance")


def test_export_remains_separate_from_recovery_and_refreshes_after_restore(process):
    client, _, *_ = process
    first = export(client)
    old_identity = headers(client)
    saved = backup(client)
    create(client, "products", name="Added after recovery point")
    preview = inspect(client, saved)
    restore(client, saved, preview)
    status = client.get(PATH).json()
    assert status["spreadsheet"]["generation"] != status["generation"]
    export(client, expected=409, request_headers=old_identity)
    second = export(client)
    assert second["id"] == first["id"] and second["generation"] == status["generation"]
    assert second["counts"] == first["counts"]
    book, _ = download(client)
    assert "Added after recovery point" not in {row["Product"] for row in rows(book, "Products")}


def test_invalid_or_oversized_report_never_replaces_existing_workbook(process):
    client, settings, product, *_ = process
    export(client)
    _, before = download(client)
    with patch("backend.app.manufacturing_spreadsheet.MAX_CELLS", 5):
        export(client, expected=413)
    file = next((settings.data_dir / "producttrackr-spreadsheets").rglob("*.xlsx"))
    assert file.read_bytes() == before
    repo = ManufacturingRepository(settings)
    repo.ref("products", product["id"]).update({"stepIds": ["missing-step"]})
    export(client, expected=422)
    assert file.read_bytes() == before
    repo.ref("products", product["id"]).update({"stepIds": []})
    for item in repo.collection("assignments").stream():
        item.reference.delete()
    for item in repo.collection("steps").stream():
        item.reference.delete()
    repo.ref("products", product["id"]).update({"notes": "bad\x01value"})
    export(client, expected=422)
    assert file.read_bytes() == before


def mock_drive():
    store = object.__new__(DriveSpreadsheetStore)
    store.scope, store.folder_id = "workspace-scope", "producttrackr-folder"
    store._folder = Mock(return_value=store.folder_id)
    store.backend = Mock()
    store.backend.etag.return_value = '"workbook-etag"'
    store.backend._execute.side_effect = lambda request: request.execute()
    store.backend._conditional_api.return_value.files.return_value.update.return_value.headers = {}
    files = store.backend._api.return_value.files.return_value
    files.get.return_value.execute.return_value = {
        "id": "workbook-id",
        "parents": [store.folder_id],
        "mimeType": MIME,
        "appProperties": {"app": store.app, "workspace": store.scope},
        "size": "10",
    }
    store.backend._api.return_value.permissions.return_value.list.return_value.execute.return_value = {
        "permissions": [{"type": "user"}],
    }
    return store, files


def test_drive_reuses_scoped_workbook_and_updates_only_its_media():
    store, files = mock_drive()
    store.backend._list.return_value = [{"id": "workbook-id"}]
    assert store.allocate() == "workbook-id"
    assert "workspace-scope" in store.backend._list.call_args.args[0]
    store.put("workbook-id", b"new workbook", existing=True)
    updates = store.backend._conditional_api.return_value.files.return_value.update
    assert updates.call_args.kwargs["fileId"] == "workbook-id"
    assert updates.call_args.kwargs["media_body"].getbytes(0, -1) == b"new workbook"
    assert updates.return_value.headers == {"If-Match": '"workbook-etag"'}
    files.create.assert_not_called()
    store.backend._list.return_value = [{"id": "first"}, {"id": "second"}]
    with pytest.raises(Problem, match="More than one"):
        store.allocate()
    files.get.return_value.execute.side_effect = Problem("Missing", 503, "drive_missing")
    with pytest.raises(Problem, match="missing"):
        store.put("workbook-id", b"new workbook", existing=True)
    files.create.assert_not_called()
    store.put("preallocated-id", b"first workbook")
    assert files.create.call_args.kwargs["body"] == {
        "id": "preallocated-id",
        "name": NAME,
        "mimeType": MIME,
        "parents": [store.folder_id],
        "appProperties": {"app": store.app, "workspace": store.scope},
    }


@pytest.mark.parametrize(
    "changed",
    [
        {"parents": ["other-folder"]},
        {"trashed": True},
        {"mimeType": "application/vnd.google-apps.spreadsheet"},
        {"appProperties": {"app": "unrelated", "workspace": "workspace-scope"}},
        {"appProperties": {"app": "producttrackr-directory-spreadsheet", "workspace": "other"}},
    ],
)
def test_drive_rejects_moved_or_foreign_workbook(changed):
    store, files = mock_drive()
    files.get.return_value.execute.return_value.update(changed)
    with pytest.raises(Problem, match="moved or replaced"):
        store.put("workbook-id", b"must not upload", existing=True)
    files.update.assert_not_called()
    store.backend._conditional_api.return_value.files.return_value.update.assert_not_called()


def test_drive_rejects_public_workbook_without_changing_permissions():
    store, files = mock_drive()
    permissions = store.backend._api.return_value.permissions.return_value
    permissions.list.return_value.execute.return_value = {"permissions": [{"type": "anyone"}]}
    with pytest.raises(Problem, match="private"):
        store.put("workbook-id", b"must not upload", existing=True)
    files.update.assert_not_called()
    store.backend._conditional_api.return_value.files.return_value.update.assert_not_called()
    permissions.create.assert_not_called()
    permissions.update.assert_not_called()


def test_drive_rechecks_ownership_after_preparing_conditional_upload():
    store, _ = mock_drive()
    callback = Mock(side_effect=Problem("Export ownership changed", 409, "workspace_changed"))
    with pytest.raises(Problem, match="ownership changed"):
        store.put("workbook-id", b"must not upload", existing=True, before_write=callback)
    callback.assert_called_once()
    store.backend._conditional_api.return_value.files.return_value.update.return_value.execute.assert_not_called()


def test_drive_mode_never_silently_falls_back_to_local_storage(app):
    client, settings = app
    with pytest.raises(Problem, match="reconnect Google Drive"):
        SpreadsheetExport(
            ManufacturingRepository(settings),
            replace(
                settings,
                manufacturing_backup_storage="drive",
                root_id="configured-folder",
            ),
        )
    assert client.get(PATH).json()["provider"] == "local"
